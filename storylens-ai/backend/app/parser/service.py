"""
Document parsing service.

Extracts plain text from various file types:
- PDF   → PyMuPDF (fitz)
- DOCX  → python-docx
- PPTX  → python-pptx
- TXT   → raw text
- Images → Tesseract OCR

Also includes a chunking utility used to prepare content for vector search.
"""
import logging
import os
from pathlib import Path
from typing import List, Tuple

from app.ocr.service import ocr_service

logger = logging.getLogger(__name__)


class DocumentParser:
    """Extracts and chunks text from uploaded documents."""

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".pptx",
        ".txt",
        ".md",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".tiff",
        ".bmp",
    }

    # ------------------------------------------------------------------
    def extract_text(self, file_path: str) -> str:
        """Dispatch to the appropriate extractor based on file extension."""
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            return self._extract_pdf(file_path)
        if ext == ".docx":
            return self._extract_docx(file_path)
        if ext == ".pptx":
            return self._extract_pptx(file_path)
        if ext in {".txt", ".md"}:
            return self._extract_text_file(file_path)
        if ext in ocr_service.SUPPORTED_EXTENSIONS:
            return ocr_service.extract_text(file_path)
        raise ValueError(f"Unsupported file type: {ext}")

    # ------------------------------------------------------------------
    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from a PDF using PyMuPDF."""
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("PyMuPDF is not installed.") from exc

        pages: List[str] = []
        with fitz.open(file_path) as doc:
            for page in doc:
                pages.append(page.get_text())
        return "\n\n".join(pages).strip()

    # ------------------------------------------------------------------
    def _extract_docx(self, file_path: str) -> str:
        """Extract text from a Word document using python-docx."""
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is not installed.") from exc

        doc = DocxDocument(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Include table text as well
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts).strip()

    # ------------------------------------------------------------------
    def _extract_pptx(self, file_path: str) -> str:
        """Extract text from a PowerPoint file using python-pptx."""
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-pptx is not installed.") from exc

        prs = Presentation(file_path)
        parts: List[str] = []
        for idx, slide in enumerate(prs.slides, start=1):
            slide_parts: List[str] = [f"[Slide {idx}]"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs).strip()
                        if text:
                            slide_parts.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        ]
                        if cells:
                            slide_parts.append(" | ".join(cells))
            parts.append("\n".join(slide_parts))
        return "\n\n".join(parts).strip()

    # ------------------------------------------------------------------
    @staticmethod
    def _extract_text_file(file_path: str) -> str:
        """Read a plain-text file with encoding fallbacks."""
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                with open(file_path, "r", encoding=encoding) as fh:
                    return fh.read().strip()
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode text file.")

    # ------------------------------------------------------------------
    @staticmethod
    def chunk_text(
        text: str, chunk_size: int = 1200, overlap: int = 150
    ) -> List[str]:
        """
        Split text into overlapping semantic chunks.

        Uses LangChain's RecursiveCharacterTextSplitter when available,
        with a simple fallback otherwise.
        """
        try:
            from langchain_text_splitters import RecursiveCharacterTextSplitter

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=chunk_size,
                chunk_overlap=overlap,
                separators=["\n\n", "\n", ". ", " ", ""],
            )
            return splitter.split_text(text)
        except ImportError:  # pragma: no cover
            return _simple_chunk(text, chunk_size, overlap)


def _simple_chunk(text: str, chunk_size: int, overlap: int) -> List[str]:
    """Minimal whitespace-based chunker (fallback)."""
    words = text.split()
    chunks: List[str] = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end])
        if chunk:
            chunks.append(chunk)
        if end == len(words):
            break
        start = max(end - overlap, start + 1)
    return chunks


document_parser = DocumentParser()

